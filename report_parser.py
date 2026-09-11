import os
import re
import unicodedata
import csv
import io
import uuid
import hashlib
import zipfile
import xml.etree.ElementTree as ET
from typing import Optional, List

from openpyxl import load_workbook
from docx import Document
from pypdf import PdfReader
from openai import OpenAI
from pydantic import BaseModel, Field



# ============================================================
# MODELOS ESTRUCTURADOS DE IA
# ============================================================

class AuditProposal(BaseModel):
    proposal_text: str = ""
    source_supported: bool = True


class AuditAnalysis(BaseModel):
    title: str
    situation: str
    evidence: str = ""
    affected_process_or_control: str = ""
    cause: Optional[str] = None
    risk: str = ""
    impact: Optional[str] = None
    severity: str = "Medio"
    responsible_area: Optional[str] = None
    proposals: List[AuditProposal] = Field(default_factory=list)
    confidence: float = 0.0


class AuditReview(BaseModel):
    approved: bool
    score: float
    issues: List[str] = Field(default_factory=list)
    corrected_result: Optional[AuditAnalysis] = None


# ============================================================
# UTILIDADES GENERALES
# ============================================================

def normalize_text(value):
    if value is None:
        return ""

    text = str(value).strip()

    if not text:
        return ""

    text = unicodedata.normalize("NFD", text)
    text = "".join(
        ch for ch in text
        if unicodedata.category(ch) != "Mn"
    )

    text = re.sub(r"\s+", " ", text.lower())

    return text.strip()


def clean_text(value):
    if value is None:
        return ""

    return re.sub(
        r"\s+",
        " ",
        str(value)
    ).strip()


# ============================================================
# CLIENTE OPENAI
# ============================================================

def _get_openai_client():
    api_key = os.getenv("OPENAI_API_KEY")

    if not api_key:
        print("[IA] OPENAI_API_KEY no configurada. Se usará fallback heurístico.")
        return None

    return OpenAI(api_key=api_key)


# ============================================================
# EXTRACCIÓN DE METADATOS
# ============================================================

def extract_explicit_area(raw_text):
    if not raw_text:
        return "Operaciones"

    patterns = [
        r"(?:área auditada|area auditada)\s*[:\-]\s*([^\n\r\|]+)",
        r"(?:proceso auditado|proceso)\s*[:\-]\s*([^\n\r\|]+)",
        r"(?:sector)\s*[:\-]\s*([^\n\r\|]+)",
        r"(?:gerencia)\s*[:\-]\s*([^\n\r\|]+)",
        r"(?:departamento)\s*[:\-]\s*([^\n\r\|]+)",
        r"(?:unidad auditada)\s*[:\-]\s*([^\n\r\|]+)",
        r"(?:alcance)\s*[:\-]\s*([^\n\r\|]+)",
    ]

    for pattern in patterns:
        match = re.search(
            pattern,
            raw_text,
            re.IGNORECASE
        )

        if match:
            result = clean_text(match.group(1))

            if 3 < len(result) < 80:
                return result

    return "Operaciones"


def extract_explicit_auditor(raw_text):
    if not raw_text:
        return "Auditoría Interna"

    pattern = (
        r"(?:auditor responsable|auditor líder|auditor lider|"
        r"auditor encargado|auditor|elaborado por|realizado por)"
        r"\s*[:\-]\s*([^\n\r\|]+)"
    )

    match = re.search(
        pattern,
        raw_text,
        re.IGNORECASE
    )

    if match:
        result = clean_text(match.group(1))

        if 2 < len(result) < 80:
            return result

    return "Auditoría Interna"


# ============================================================
# EXTRACCIÓN DE TEXTO DE ARCHIVOS
# ============================================================

def extract_docx_xml_deep(file_path):
    text_parts = []
    try:
        with zipfile.ZipFile(file_path, 'r') as z:
            for name in z.namelist():
                if name.startswith("word/") and name.endswith(".xml"):
                    try:
                        xml_bytes = z.read(name)
                        root = ET.fromstring(xml_bytes)
                        for elem in root.iter():
                            if elem.tag.endswith('}t') and elem.text:
                                text_parts.append(elem.text.strip())
                            elif elem.tag.endswith('}p') or elem.tag.endswith('}tr'):
                                text_parts.append("\n")
                    except Exception:
                        pass
    except Exception as e:
        print(f"[Parser] Excepción en extracción XML de ZIP: {e}")

    raw_res = " ".join(text_parts)
    lines = [clean_text(l) for l in raw_res.splitlines() if clean_text(l)]
    return "\n".join(lines)


def extract_raw_text_from_file(file_path):
    ext = (
        file_path.rsplit(".", 1)[-1].lower()
        if "." in file_path
        else ""
    )

    text_content = ""

    if ext in ("docx", "doc"):
        # Tier 1: python-docx (párrafos, tablas y encabezados)
        try:
            doc = Document(file_path)

            paragraphs = [
                p.text.strip()
                for p in doc.paragraphs
                if p.text.strip()
            ]

            for table in doc.tables:
                for row in table.rows:
                    row_cells = [
                        c.text.strip()
                        for c in row.cells
                        if c.text.strip()
                    ]

                    if row_cells:
                        paragraphs.append(
                            " | ".join(row_cells)
                        )

            for section in doc.sections:
                if section.header:
                    hdr_txt = "\n".join([p.text.strip() for p in section.header.paragraphs if p.text.strip()])
                    if hdr_txt:
                        paragraphs.insert(0, f"[Encabezado: {hdr_txt}]")

            text_content = "\n".join(paragraphs)

        except Exception as exc:
            print(f"[Parser] python-docx aviso en {file_path}: {exc}")

        # Tier 2: Extracción profunda XML del archivo ZIP (cuadros de texto, formas, marcos)
        if not text_content.strip():
            text_content = extract_docx_xml_deep(file_path)

        # Tier 3: Fallback de decodificación directa de cadenas para archivos binarios (.doc)
        if not text_content.strip():
            try:
                with open(file_path, "rb") as f_raw:
                    raw_b = f_raw.read()
                decoded = raw_b.decode("latin-1", errors="ignore")
                printable = re.findall(r'[A-Za-z0-9ÁÉÍÓÚáéíóúÑñ\s\.,;:!\?\-\(\)\$/]{4,}', decoded)
                text_content = "\n".join([p.strip() for p in printable if len(p.strip()) > 5])
            except Exception:
                pass

    elif ext == "pdf":
        try:
            reader = PdfReader(file_path)
            pages_text = []
            total_pages = len(reader.pages)

            for idx, page in enumerate(reader.pages):
                txt = page.extract_text() or ""

                if txt.strip():
                    pages_text.append(
                        f"--- PÁGINA {idx + 1} ---\n{txt}"
                    )

            text_content = "\n".join(pages_text)

            if total_pages > 0 and not text_content.strip():
                raise ValueError("El archivo PDF no contiene texto seleccionable o es una imagen escaneada (requiere OCR).")

        except ValueError:
            raise
        except Exception as exc:
            raise ValueError(f"No se pudo procesar el archivo PDF: {str(exc)}")

    elif ext == "xlsx":
        try:
            wb = load_workbook(
                file_path,
                data_only=True
            )

            lines = []
            max_sheets = 10

            for sheet_name in wb.sheetnames[:max_sheets]:
                ws = wb[sheet_name]

                lines.append(
                    f"=== SOLAPA: {sheet_name} ==="
                )

                row_count = 0
                for row in ws.iter_rows(
                    values_only=True
                ):
                    row_count += 1
                    if row_count > 2000:
                        lines.append("... [Límite de 2000 filas alcanzado en esta solapa] ...")
                        break

                    row_vals = [
                        clean_text(v)
                        for v in row
                        if clean_text(v)
                    ]

                    if row_vals:
                        lines.append(
                            " | ".join(row_vals)
                        )

            text_content = "\n".join(lines)

            if not text_content.strip():
                wb_fallback = load_workbook(file_path, data_only=False)
                lines_fb = []
                for sheet_name in wb_fallback.sheetnames[:max_sheets]:
                    ws = wb_fallback[sheet_name]
                    lines_fb.append(f"=== SOLAPA: {sheet_name} ===")
                    for row in ws.iter_rows(values_only=True):
                        row_vals = [clean_text(v) for v in row if clean_text(v)]
                        if row_vals:
                            lines_fb.append(" | ".join(row_vals))
                text_content = "\n".join(lines_fb)

        except Exception as exc:
            raise ValueError(f"No se pudo procesar la planilla Excel: {str(exc)}")

    elif ext == "csv":
        try:
            with open(file_path, "rb") as f_raw:
                raw_bytes = f_raw.read()

            try:
                content_str = raw_bytes.decode("utf-8-sig")
            except UnicodeDecodeError:
                content_str = raw_bytes.decode("latin-1", errors="ignore")

            sample = content_str[:4096]
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=[',', ';', '\t', '|'])
                delimiter = dialect.delimiter
            except Exception:
                delimiter = ';' if ';' in sample else ','

            reader = csv.reader(io.StringIO(content_str), delimiter=delimiter)
            rows_str = []
            for row in reader:
                cleaned_cells = [clean_text(c) for c in row if clean_text(c)]
                if cleaned_cells:
                    rows_str.append(" | ".join(cleaned_cells))

            text_content = "\n".join(rows_str)

        except Exception as exc:
            raise ValueError(f"No se pudo procesar el archivo CSV: {str(exc)}")

    elif ext == "txt":
        try:
            with open(file_path, "rb") as f_raw:
                raw_bytes = f_raw.read()

            try:
                text_content = raw_bytes.decode("utf-8-sig")
            except UnicodeDecodeError:
                text_content = raw_bytes.decode("latin-1", errors="ignore")

        except Exception as exc:
            raise ValueError(f"No se pudo procesar el archivo TXT: {str(exc)}")

    if not text_content or not text_content.strip():
        raise ValueError("El archivo no contiene texto interpretable o está vacío.")

    return text_content[:60000]


# ============================================================
# PARSER HEURÍSTICO - EXTRACTOR / FALLBACK
# ============================================================

NON_FINDING_KEYWORDS = {
    "alcance",
    "metodologia",
    "metodología",
    "objetivo",
    "objetivos",
    "introduccion",
    "introducción",
    "antecedentes",
    "contexto",
    "conclusion",
    "conclusión",
    "conclusiones",
    "resumen ejecutivo",
    "anexo",
    "anexos",
    "referencias",
    "glosario",
    "índice",
    "indice",
    "marco normativo",
    "marco de referencia",
    "cronograma",
    "equipo auditor",
    "distribucion",
    "distribución",
    "carátula",
    "caratula",
    "portada",
    "periodo",
    "período",
    "tabla de contenido",
    "contenido",
    "contenidos",
}


FINDING_KEYWORDS = [
    "hallazgo",
    "observacion",
    "observación",
    "desviacion",
    "desviación",
    "debilidad",
    "deficiencia",
    "incumplimiento",
    "punto de atención",
    "punto de atencion",
    "irregularidad",
]


PROPOSAL_KEYWORDS = [
    "propuesta de mejora",
    "propuesta",
    "recomendacion",
    "recomendación",
    "accion correctiva",
    "acción correctiva",
    "mejora sugerida",
    "sugerencia",
    "medida correctiva",
    "accion recomendada",
    "acción recomendada",
]


def _is_non_finding_header(norm_line):
    for kw in NON_FINDING_KEYWORDS:
        pattern = (
            r"(?:^|\d+[\.\s]+)"
            + re.escape(kw)
            + r"s?\s*[:.]?\s*$"
        )

        if re.search(
            pattern,
            norm_line
        ):
            return True

        if (
            norm_line
            .strip()
            .rstrip(":. ")
            == kw
        ):
            return True

    return False


PROPOSAL_VERBS = (
    "reforzar",
    "formalizar",
    "definir",
    "actualizar",
    "incorporar",
    "implementar",
    "establecer",
    "asegurar",
    "revisar",
)

EXCLUDED_HEADERS = [
    "categoría de control",
    "categoria de control",
    "riesgo detectado",
    "descripción del problema",
    "descripcion del problema",
    "impacto y consecuencias",
    "propuesta de mejora",
    "propuesta",
    "recomendación",
    "recomendacion",
    "medida correctiva",
    "acción correctiva",
    "accion correctiva",
]


def _looks_like_finding_start(
    line,
    norm_line
):
    if "|" in line or "[table_row]" in norm_line or "[table" in norm_line:
        return False

    explicit_re = re.compile(
        r"^(?:\d+[\.\)]\s*)?(?:hallazgo|observaci[oó]n|"
        r"desviaci[oó]n|punto|ítem|item)"
        r"\s*(?:n[°º]?\s*)?\d+",
        re.IGNORECASE
    )

    if explicit_re.match(line.strip()):
        return True

    hallazgo_header_re = re.compile(
        r"^(?:\d+[\.\)]\s*)?hallazgo\b",
        re.IGNORECASE
    )

    if len(line.strip()) < 120 and hallazgo_header_re.match(line.strip()):
        return True

    for verb in PROPOSAL_VERBS:
        if norm_line.startswith(verb) or re.match(r"^(?:\d+[\.\)]\s*)?" + verb + r"\b", norm_line):
            return False

    for header in EXCLUDED_HEADERS:
        if norm_line.startswith(header) or norm_line.strip().rstrip(":. ") == header:
            return False

def _is_proposals_section_header(line, norm_line):
    clean_l = re.sub(r"\[[A-Z0-9_]+\]", "", line).strip()
    header_text = re.sub(r"^(?:secci[oó]n|cap[ií]tulo)?\s*(?:\d+[\.\:\)\-]*|\d+\.\d+[\.\:\)\-]*|[-•])\s*", "", clean_l, flags=re.IGNORECASE).strip()
    header_text = header_text.rstrip(":. -")
    norm_h = normalize_text(header_text)

    if norm_h in (
        "propuestas de mejora",
        "propuesta de mejora",
        "propuestas",
        "propuesta de mejoras",
        "recomendaciones de mejora",
        "recomendaciones",
        "acciones correctivas"
    ):
        return True

    if any(kw in norm_h for kw in ["propuesta", "recomendacion", "recomendación", "accion correctiva", "acción correctiva"]):
        if any(w in norm_h for w in ["mejora", "propuestas", "recomendaciones", "acciones"]):
            return True

    return False


def _parse_finding_references(text):
    if not text:
        return []

    found_numbers = set()
    pattern = re.compile(
        r"\b(?:hallazgos?|observaci[oó]n(?:es)?|desviaci[oó]n(?:es)?|h)\s*(?:n[°ºº]?\s*|-)?(\d+(?:\s*(?:,|y|e)\s*\d+)*)",
        re.IGNORECASE
    )

    for match in pattern.finditer(text):
        nums_str = match.group(1)
        nums = re.findall(r"\d+", nums_str)
        for n in nums:
            found_numbers.add(int(n))

    return sorted(list(found_numbers))


def _looks_like_proposal_start(
    line,
    norm_line
):
    for kw in PROPOSAL_KEYWORDS:
        if norm_line.startswith(kw):
            return True

        pattern = (
            re.escape(kw)
            + r"\s*[:\-]"
        )

        if re.match(
            pattern,
            norm_line
        ):
            return True

    return False


def _extract_after_keyword(
    line
):
    for kw in PROPOSAL_KEYWORDS:
        pattern = re.compile(
            re.escape(kw)
            + r"\s*[:\-]\s*",
            re.IGNORECASE
        )

        match = pattern.match(line)

        if match:
            return line[
                match.end():
            ].strip()

    return line


def _detect_severity(text):
    if not text:
        return "Medio"

    norm = normalize_text(text)

    # 1. Explicit risk headers and phrases
    if any(phrase in norm for phrase in [
        "riesgo alto", "riesgo: alto", "riesgo alta", "riesgo: alta",
        "criticidad alta", "criticidad: alta", "severidad alta", "severidad: alta",
        "riesgo critico", "riesgo crítico", "criticidad critica", "criticidad crítica",
        "riesgo  alto", "riesgo   alto"
    ]):
        return "Alto"

    if any(phrase in norm for phrase in [
        "riesgo bajo", "riesgo: bajo", "riesgo baja", "riesgo: baja",
        "criticidad baja", "criticidad: baja", "severidad baja", "severidad: baja",
        "riesgo leve", "riesgo menor"
    ]):
        return "Bajo"

    if any(phrase in norm for phrase in [
        "riesgo medio", "riesgo: medio", "riesgo media", "riesgo: media",
        "criticidad media", "criticidad: media", "severidad media", "severidad: media",
        "riesgo moderado", "riesgo moderada"
    ]):
        return "Medio"

    # 2. General word matching
    if any(
        word in norm
        for word in [
            "alto",
            "alta",
            "critico",
            "critica",
            "crítico",
            "crítica",
        ]
    ):
        return "Alto"

    if any(
        word in norm
        for word in [
            "bajo",
            "baja",
            "leve",
            "menor",
        ]
    ):
        return "Bajo"

    return "Medio"


def _clean_finding_title(
    raw_title
):
    title = clean_text(raw_title)

    title = re.sub(
        r"^(?:hallazgo|observaci[oó]n|"
        r"desviaci[oó]n|punto|ítem|item)"
        r"\s*(?:n[°º]?\s*)?\d*"
        r"\s*[:\-\.]\s*",
        "",
        title,
        flags=re.IGNORECASE
    ).strip()

    if not title:
        return clean_text(raw_title)[:100]

    if len(title) > 100:
        cut = title[:100].rfind(" ")

        title = (
            title[:cut]
            if cut > 25
            else title[:100]
        )

    if title:
        title = (
            title[0].upper()
            + title[1:]
        )

    return title


def _summarize_text(
    text,
    max_chars=700
):
    text = clean_text(text)

    if not text:
        return ""

    if len(text) <= max_chars:
        return text

    sentences = re.split(
        r"(?<=[.!?])\s+",
        text
    )

    selected = []
    total = 0

    for sentence in sentences:
        sentence = sentence.strip()

        if not sentence:
            continue

        if (
            total
            + len(sentence)
            + 1
            > max_chars
        ):
            break

        selected.append(sentence)

        total += (
            len(sentence)
            + 1
        )

    if selected:
        return " ".join(selected)

    return text[:max_chars]


# ============================================================
# IA CAPA 1 - ANALISTA SENIOR
# ============================================================

def analyze_finding_with_ai(
    source_text,
    fallback_title="",
    fallback_area="",
    fallback_severity="Medio",
):
    client = _get_openai_client()

    if not client:
        return None

    source_text = source_text.strip()

    if not source_text:
        return None

    instructions = """
Actuás como Auditor Interno Senior.

Tu función es analizar información proveniente de informes de Auditoría
Interna de cualquier proceso, área o temática.

No asumas que el informe corresponde a inventarios, stock, sistemas,
contabilidad, legales, proveedores ni ninguna temática específica.

Debés interpretar exclusivamente el contenido proporcionado.

OBJETIVO

Identificar el verdadero hallazgo de auditoría y estructurarlo de forma
profesional.

Cuando la información esté disponible, distinguí:

- situación detectada;
- evidencia concreta;
- proceso o control afectado;
- causa;
- riesgo o consecuencia;
- impacto cuantitativo;
- severidad;
- área responsable;
- propuesta o propuestas de mejora.

REGLAS PARA "SITUACIÓN OBSERVADA" (campo 'situation'):

1. Debe ser una SÍNTESIS EJECUTIVA Y RESUMIDA DE MÁXIMO ~500 CARACTERES.
2. Debe explicar sintéticamente:
   - qué ocurrió (hecho o desvío clave);
   - cuál fue la debilidad o incumplimiento identificado;
   - cuál es la consecuencia o control afectado cuando corresponda.
3. RESTRICCIONES ESTRICTAS:
   - NO copiar textualmente todo el análisis o relato detallado del informe.
   - NO copiar tablas de datos ni listados de items.
   - NO repetir conversaciones o entrevistas ("Se consultó a Stock...", "El área indicó...", "En reunión se acordó...").
   - NO relatar paso a paso la investigación o el trabajo de auditoría realizado.
   - NO inventar hechos, cifras o información que no figure en la fuente.
   - NO copiar íntegramente la sección de conclusión ni la concatenación de párrafos.

REGLAS

1. NO INVENTAR INFORMACIÓN.

No inventes:
- importes;
- porcentajes;
- cantidades;
- fechas;
- códigos;
- documentos;
- transacciones;
- sistemas;
- áreas;
- responsables;
- normativa;
- causas;
- hechos.

2. Si una causa no está respaldada por el texto fuente, devolver null.

3. Si no existe impacto cuantitativo, devolver null.

4. Conservá la precisión de:
- cifras;
- fechas;
- códigos;
- sistemas;
- documentos;
- procesos;
- evidencia.

5. No confundas comentarios de reuniones, respuestas del área o frases
operativas con el hallazgo principal.

6. Expresiones aisladas como:
"falta de control",
"se detectaron diferencias",
"error de sistema",
"posible incumplimiento",
"falta de seguimiento"
no constituyen por sí mismas un hallazgo completo.

7. El hallazgo debe explicar claramente:
qué situación fue identificada y por qué resulta relevante desde la
perspectiva de Auditoría Interna.

8. El riesgo debe explicar la consecuencia razonable de la situación
identificada, sin exageraciones.

9. No usar riesgos genéricos tipo:
"riesgo de control interno asociado al hallazgo".

10. La propuesta debe responder directamente al problema detectado.

11. Evitar propuestas genéricas como:
"mejorar controles",
"realizar seguimiento",
"ejecutar y documentar",
"capacitar al personal",
salvo que exista sustento concreto para esa medida.

12. Si existe una recomendación original en el documento, preservá su sentido.

13. Puede haber más de una propuesta para un mismo hallazgo.

14. Hallazgo y propuesta no deben ser una paráfrasis uno del otro.

15. Severidad solamente puede ser:
Alto, Medio o Bajo.

16. Si el documento no permite determinar la severidad claramente,
usar Medio.

17. La redacción debe ser técnica, profesional, concreta y natural.

18. No agregar un plan de acción. Las propuestas de mejora y los planes
de acción son conceptos diferentes.
"""

    try:
        user_prompt = f"""
Analizá este bloque de un informe de Auditoría Interna.

TÍTULO PRELIMINAR DETECTADO:
{fallback_title}

ÁREA PRELIMINAR:
{fallback_area}

SEVERIDAD PRELIMINAR:
{fallback_severity}

TEXTO FUENTE:

--- INICIO TEXTO FUENTE ---

{source_text[:14000]}

--- FIN TEXTO FUENTE ---

Los valores preliminares fueron obtenidos mediante reglas heurísticas.
Usalos únicamente como referencia.

La fuente original prevalece sobre cualquier dato preliminar.
"""
        response = client.beta.chat.completions.parse(
            model=os.getenv("OPENAI_AUDIT_MODEL", "gpt-4o"),
            messages=[
                {"role": "system", "content": instructions},
                {"role": "user", "content": user_prompt}
            ],
            response_format=AuditAnalysis,
        )

        return response.choices[0].message.parsed

    except Exception as exc:
        print(f"[IA Analista] Error de API u OpenAI sin crédito: {exc}")
        return None


# ============================================================
# IA CAPA 2 - REVISOR INDEPENDIENTE
# ============================================================

def review_ai_analysis(
    source_text,
    analysis
):
    client = _get_openai_client()

    if (
        not client
        or analysis is None
    ):
        return None

    instructions = """
Actuás como Revisor Senior Independiente de Auditoría Interna.

Otro auditor IA realizó un análisis.

Tu función NO es volver a analizar libremente el caso desde cero.

Debés comparar su resultado contra la fuente original.

CONTROLAR

1. que el hallazgo esté respaldado por la fuente;
2. que no existan hechos inventados;
3. que no existan cifras inventadas;
4. que no existan fechas inventadas;
5. que no existan códigos inventados;
6. que no existan sistemas inventados;
7. que no existan áreas o responsables inventados;
8. que no se haya inventado una causa;
9. que el riesgo sea razonable;
10. que el riesgo no sea genérico;
11. que la propuesta responda al hallazgo;
12. que la propuesta no sea genérica;
13. que hallazgo y propuesta sean conceptos distintos;
14. que no se hayan omitido datos relevantes;
15. que las recomendaciones originales hayan sido preservadas;
16. que no se haya transformado una opinión o comentario en un hecho;
17. que el resultado sea claro y profesional.

SI TODO ES CORRECTO

approved = true
corrected_result = null

SI EXISTEN ERRORES

approved = false

Explicar cada problema en issues.

Generar corrected_result solamente corrigiendo los problemas detectados.

No agregar información nueva durante la corrección.
"""

    try:
        user_prompt = f"""
TEXTO FUENTE ORIGINAL

--- INICIO FUENTE ---

{source_text[:14000]}

--- FIN FUENTE ---


RESULTADO GENERADO POR LA PRIMERA IA

--- INICIO RESULTADO ---

{analysis.model_dump_json(indent=2)}

--- FIN RESULTADO ---


Revisá exclusivamente la trazabilidad, consistencia y calidad del resultado.
"""
        response = client.beta.chat.completions.parse(
            model=os.getenv("OPENAI_AUDIT_REVIEW_MODEL", "gpt-4o"),
            messages=[
                {"role": "system", "content": instructions},
                {"role": "user", "content": user_prompt}
            ],
            response_format=AuditReview,
        )

        return response.choices[0].message.parsed

    except Exception as exc:
        print(f"[IA Revisora] Error de API u OpenAI sin crédito: {exc}")
        return None


# ============================================================
# COORDINADOR DE DOBLE CONTROL
# ============================================================

def run_double_ai_review(
    source_text,
    fallback_title="",
    fallback_area="",
    fallback_severity="Medio",
):
    analysis = analyze_finding_with_ai(
        source_text=source_text,
        fallback_title=fallback_title,
        fallback_area=fallback_area,
        fallback_severity=fallback_severity,
    )

    if not analysis:
        return None

    review = review_ai_analysis(
        source_text,
        analysis
    )

    if review is None:
        print(
            "[IA] No fue posible ejecutar "
            "la segunda revisión."
        )

        return None

    if review.approved:
        return analysis

    if review.corrected_result:
        corrected = (
            review.corrected_result
        )

        second_review = review_ai_analysis(
            source_text,
            corrected
        )

        if (
            second_review
            and second_review.approved
        ):
            return corrected

    print(
        "[IA] El resultado no superó "
        "la doble revisión. "
        "Se utilizará fallback heurístico."
    )

    return None


# ============================================================
# PARSEO DE DOCUMENTO EN BLOQUES
# ============================================================

def parse_document(
    raw_text,
    filename
):
    explicit_area = (
        extract_explicit_area(raw_text)
    )

    explicit_auditor = (
        extract_explicit_auditor(raw_text)
    )

    report_title = (
        f"Informe de Auditoría - {filename}"
    )

    lines = [
        clean_text(line)
        for line in raw_text.splitlines()
        if clean_text(line)
    ]

    for line in lines[:20]:
        norm = normalize_text(line)

        if (
            "informe" in norm
            and "auditoria" in norm
        ):
            report_title = line
            break

    findings = []
    global_proposals = []
    current_finding = None
    current_global_proposal = None
    reading_proposal = False
    in_global_proposals = False
    prev_line = ""

    proposal_start_verbs = {
        "incorporar", "reforzar", "asegurar", "formalizar", "definir", "actualizar",
        "evaluar", "establecer", "implementar", "revisar", "diseñar", "desarrollar",
        "modificar", "crear", "realizar", "capacitar", "monitorear", "optimizar",
        "garantizar", "promover", "solicitar", "unificar", "ajustar", "notificar",
        "supervisar", "analizar", "efectuar", "adecuar", "difundir", "gestionar"
    }

    for line in lines:
        norm = normalize_text(line)

        if _is_proposals_section_header(line, norm):
            print(f"[Propuestas] Sección global encontrada: {line}")
            if current_finding:
                findings.append(current_finding)
                current_finding = None

            in_global_proposals = True
            reading_proposal = False
            reading_summary = False
            reading_conclusion = False
            prev_line = line
            continue

        if in_global_proposals:
            clean_l = re.sub(r"\[[A-Z0-9_]+\]", "", line).strip()
            header_text = re.sub(r"^(?:\d+[\.\)]|\d+\.\d+[\.\)]?|[-•])\s*", "", clean_l).strip()
            norm_h = normalize_text(header_text)

            if _is_non_finding_header(norm_h) or norm_h in ("conclusion", "conclusiones", "anexo", "anexos", "glosario", "referencias", "firmas"):
                print(f"[Propuestas] Fin de sección global por encabezado: {line}")
                if current_global_proposal:
                    global_proposals.append(current_global_proposal)
                    current_global_proposal = None
                in_global_proposals = False
                prev_line = line
                continue

            num_match = re.match(r"^(?:\[[A-Z0-9_]+\]\s*)?(?:propuesta|recomendaci[oó]n)?\s*(?:n[°º]?\s*)?(\d+)[\.\:\)\-]*\s*(.*)", line.strip(), re.IGNORECASE)
            bullet_match = re.match(r"^(?:\[[A-Z0-9_]+\]\s*)?[-•\*]\s*(.*)", line.strip())

            words_h = norm_h.split()
            first_word = words_h[0] if words_h else ""
            is_verb_start = first_word in proposal_start_verbs or norm_h.startswith("propuesta") or norm_h.startswith("recomendacion")

            if num_match:
                p_num = int(num_match.group(1))
                p_text_init = num_match.group(2).strip()
                if current_global_proposal:
                    global_proposals.append(current_global_proposal)

                current_global_proposal = {
                    "number": p_num,
                    "proposal_text_lines": [p_text_init] if p_text_init else []
                }
                prev_line = line
                continue

            if bullet_match:
                p_text_init = bullet_match.group(1).strip()
                if current_global_proposal:
                    global_proposals.append(current_global_proposal)

                auto_num = (global_proposals[-1]["number"] + 1) if global_proposals else 1
                current_global_proposal = {
                    "number": auto_num,
                    "proposal_text_lines": [p_text_init] if p_text_init else []
                }
                prev_line = line
                continue

            if is_verb_start or current_global_proposal is None:
                if current_global_proposal:
                    global_proposals.append(current_global_proposal)

                auto_num = (global_proposals[-1]["number"] + 1) if global_proposals else 1
                current_global_proposal = {
                    "number": auto_num,
                    "proposal_text_lines": [line]
                }
                prev_line = line
                continue

            if current_global_proposal:
                current_global_proposal["proposal_text_lines"].append(line)
                prev_line = line
                continue

        if (
            _is_non_finding_header(norm)
            and len(line) < 120
        ):
            if current_finding:
                findings.append(
                    current_finding
                )

                current_finding = None
                reading_proposal = False

            prev_line = line
            continue

        if _looks_like_finding_start(
            line,
            norm
        ):
            if current_finding:
                findings.append(
                    current_finding
                )

            num_match = re.search(r"\b(?:hallazgo|observaci[oó]n|desviaci[oó]n|punto|ítem|item)\s*(?:n[°º]?\s*)?(\d+)", line, re.IGNORECASE)
            if not num_match:
                num_match = re.search(r"^\s*(\d+)[\.\)]", line)
            source_num = int(num_match.group(1)) if num_match else (len(findings) + 1)

            current_finding = {
                "raw_title": line,
                "prev_line": prev_line,
                "source_number": source_num,
                "situation_lines": [],
                "summary_lines": [],
                "conclusion_lines": [],
                "proposal_lines": [],
                "severity": (
                    _detect_severity(f"{prev_line} {line}")
                ),
                "responsible_area": (
                    explicit_area
                ),
            }

            reading_proposal = False
            reading_summary = False
            reading_conclusion = False
            prev_line = line
            continue

        if (
            current_finding
            and _looks_like_proposal_start(
                line,
                norm
            )
        ):
            proposal_text = (
                _extract_after_keyword(
                    line
                )
            )

            current_finding[
                "proposal_lines"
            ].append(
                proposal_text
            )

            reading_proposal = True
            reading_summary = False
            reading_conclusion = False
            prev_line = line
            continue

        if current_finding:
            if reading_proposal:
                current_finding[
                    "proposal_lines"
                ].append(line)
            else:
                if any(kw in norm for kw in ["resumen", "sintesis", "síntesis"]):
                    reading_summary = True
                    reading_conclusion = False
                elif any(kw in norm for kw in ["conclusion", "conclusión", "conclusiones"]):
                    reading_conclusion = True
                    reading_summary = False

                if reading_summary:
                    current_finding["summary_lines"].append(line)
                elif reading_conclusion:
                    current_finding["conclusion_lines"].append(line)

                current_finding[
                    "situation_lines"
                ].append(line)

        prev_line = line

    if current_finding:
        findings.append(
            current_finding
        )

    if current_global_proposal:
        global_proposals.append(current_global_proposal)

    if (
        not findings
        and "|" in raw_text
    ):
        findings = _parse_tabular_findings(
            lines,
            explicit_area
        )

        return {
            "report_title": report_title,
            "area": explicit_area,
            "auditor": explicit_auditor,
            "findings": findings,
            "global_proposals": global_proposals,
        }

    cleaned_findings = []

    for idx, finding in enumerate(findings, start=1):
        raw_title = finding.get(
            "raw_title",
            ""
        )

        situation_full = " ".join(
            finding.get(
                "situation_lines",
                []
            )
        )

        summary_full = " ".join(
            finding.get(
                "summary_lines",
                []
            )
        )

        conclusion_full = " ".join(
            finding.get(
                "conclusion_lines",
                []
            )
        )

        proposal_full = " ".join(
            finding.get(
                "proposal_lines",
                []
            )
        )

        title = _clean_finding_title(
            raw_title
        )

        source_parts = [
            raw_title,
            situation_full,
        ]

        if proposal_full:
            source_parts.append(
                "Propuesta / recomendación "
                "original:\n"
                + proposal_full
            )

        source_block = "\n".join(
            part
            for part in source_parts
            if part
        )

        full_finding_text = f"{finding.get('prev_line', '')}\n{raw_title}\n{situation_full}\n{summary_full}\n{conclusion_full}"
        detected_sev = _detect_severity(full_finding_text)
        if detected_sev == "Medio" and finding.get("severity") in ("Alto", "Bajo"):
            detected_sev = finding.get("severity")

        ai_result = run_double_ai_review(
            source_text=source_block,
            fallback_title=title,
            fallback_area=finding.get(
                "responsible_area",
                explicit_area
            ),
            fallback_severity=detected_sev,
        )

        if ai_result:
            proposals = [
                clean_text(
                    proposal.proposal_text
                )
                for proposal
                in ai_result.proposals
                if clean_text(
                    proposal.proposal_text
                )
            ]

            if clean_text(ai_result.situation):
                situation = _summarize_text(
                    clean_text(ai_result.situation),
                    max_chars=500
                ) or title
                situation_source = "IA"
                ai_status = "IA"
            else:
                summary_clean = clean_text(summary_full)
                conclusion_clean = clean_text(conclusion_full)

                if summary_clean:
                    fallback_type = "summary"
                    situation = _summarize_text(summary_clean, max_chars=500)
                elif conclusion_clean:
                    fallback_type = "conclusion"
                    situation = _summarize_text(conclusion_clean, max_chars=500)
                else:
                    fallback_type = "title"
                    situation = _summarize_text(title, max_chars=500)

                if not situation:
                    situation = title

                situation_source = f"Fallback ({fallback_type})"
                ai_status = f"Fallback ({fallback_type})"

            print(f"[Situación H{idx}] Origen: {situation_source} | Longitud: {len(situation)}")

            cleaned_findings.append({
                "title": (
                    clean_text(
                        ai_result.title
                    )
                    or title
                ),

                "source_number": finding.get(
                    "source_number",
                    idx
                ),

                "situation": situation,

                "situation_source": situation_source,

                "proposal": (
                    proposals[0]
                    if proposals
                    else ""
                ),

                "proposals_ai": proposals,

                "risk": clean_text(
                    ai_result.risk
                ),

                "severity": (
                    ai_result.severity
                    if ai_result.severity
                    in (
                        "Alto",
                        "Medio",
                        "Bajo"
                    )
                    else detected_sev
                ),

                "responsible_area": (
                    clean_text(
                        ai_result
                        .responsible_area
                    )
                    or finding.get(
                        "responsible_area",
                        explicit_area
                    )
                ),

                "evidence": clean_text(
                    ai_result.evidence
                ),

                "cause": clean_text(
                    ai_result.cause
                ),

                "affected_process_or_control":
                    clean_text(
                        ai_result
                        .affected_process_or_control
                    ),

                "impact": clean_text(
                    ai_result.impact
                ),

                "ai_confidence": (
                    ai_result.confidence
                ),

                "ai_validated": True,
                "ai_status": ai_status,

                "source_text": source_block,
            })

        else:
            summary_clean = clean_text(summary_full)
            conclusion_clean = clean_text(conclusion_full)

            if summary_clean:
                fallback_type = "summary"
                fallback_situation = _summarize_text(summary_clean, max_chars=500)
            elif conclusion_clean:
                fallback_type = "conclusion"
                fallback_situation = _summarize_text(conclusion_clean, max_chars=500)
            else:
                fallback_type = "title"
                fallback_situation = _summarize_text(title, max_chars=500)

            if not fallback_situation:
                fallback_situation = title

            situation_source = f"Fallback ({fallback_type})"
            ai_status = f"Fallback ({fallback_type})"

            print(f"[Situación H{idx}] Origen: {situation_source} | Longitud: {len(fallback_situation)}")

            fallback_proposal = (
                _summarize_text(
                    proposal_full,
                    max_chars=500
                )
            )

            cleaned_findings.append({
                "title": title,
                "source_number": finding.get(
                    "source_number",
                    idx
                ),
                "situation": (
                    fallback_situation
                ),
                "situation_source": situation_source,
                "proposal": (
                    fallback_proposal
                ),
                "proposals_ai": (
                    [fallback_proposal]
                    if fallback_proposal
                    else []
                ),
                "risk": "",
                "severity": detected_sev,
                "responsible_area": (
                    finding.get(
                        "responsible_area",
                        explicit_area
                    )
                ),
                "evidence": "",
                "cause": "",
                "affected_process_or_control": "",
                "impact": "",
                "ai_confidence": 0,
                "ai_validated": False,
                "ai_status": ai_status,
                "source_text": source_block,
            })

    print(
        f"[Parser] Documento '{filename}': "
        f"{len(cleaned_findings)} hallazgos."
    )

    return {
        "report_title": report_title,
        "area": explicit_area,
        "auditor": explicit_auditor,
        "findings": cleaned_findings,
        "global_proposals": global_proposals,
    }


# ============================================================
# PARSER TABULAR
# ============================================================

def _parse_tabular_findings(
    lines,
    default_area
):
    findings = []

    header_idx = None
    col_hallazgo = None
    col_propuesta = None
    col_area = None
    col_riesgo = None

    for index, line in enumerate(lines):
        if "|" not in line:
            continue

        cols = [
            cell.strip()
            for cell
            in line.split("|")
        ]

        norm_cols = [
            normalize_text(cell)
            for cell in cols
        ]

        if header_idx is None:
            for col_idx, norm_col in enumerate(
                norm_cols
            ):
                if any(
                    kw in norm_col
                    for kw in [
                        "hallazgo",
                        "observacion",
                        "desviacion",
                        "descripcion",
                        "situacion",
                    ]
                ):
                    col_hallazgo = col_idx

                if any(
                    kw in norm_col
                    for kw in [
                        "propuesta",
                        "recomendacion",
                        "mejora",
                    ]
                ):
                    col_propuesta = col_idx

                if any(
                    kw in norm_col
                    for kw in [
                        "area",
                        "proceso",
                        "sector",
                        "gerencia",
                    ]
                ):
                    col_area = col_idx

                if any(
                    kw in norm_col
                    for kw in [
                        "riesgo",
                        "severidad",
                        "criticidad",
                    ]
                ):
                    col_riesgo = col_idx

            if col_hallazgo is not None:
                header_idx = index
                continue

        if (
            header_idx is not None
            and col_hallazgo is not None
            and len(cols) > col_hallazgo
        ):
            hallazgo_text = (
                cols[col_hallazgo]
            )

            if len(
                hallazgo_text
            ) < 5:
                continue

            proposal_text = ""

            if (
                col_propuesta
                is not None
                and len(cols)
                > col_propuesta
            ):
                proposal_text = (
                    cols[col_propuesta]
                )

            area = default_area

            if (
                col_area
                is not None
                and len(cols)
                > col_area
            ):
                area = (
                    cols[col_area]
                    or default_area
                )

            severity = "Medio"

            if (
                col_riesgo
                is not None
                and len(cols)
                > col_riesgo
            ):
                severity = (
                    _detect_severity(
                        cols[col_riesgo]
                    )
                )

            source_block = (
                hallazgo_text
            )

            if proposal_text:
                source_block += (
                    "\nPropuesta / recomendación "
                    "original:\n"
                    + proposal_text
                )

            ai_result = run_double_ai_review(
                source_text=source_block,
                fallback_title=hallazgo_text[:100],
                fallback_area=area,
                fallback_severity=severity,
            )

            situation = ""
            situation_source = ""
            ai_status = ""

            if ai_result and clean_text(ai_result.situation):
                situation = _summarize_text(clean_text(ai_result.situation), max_chars=500)
                situation_source = "IA"
                ai_status = "IA"
            else:
                hallazgo_clean = clean_text(hallazgo_text)
                if hallazgo_clean:
                    fallback_type = "summary"
                    situation = _summarize_text(hallazgo_clean, max_chars=500)
                else:
                    fallback_type = "title"
                    situation = _summarize_text(hallazgo_text[:100], max_chars=500)
                situation_source = f"Fallback ({fallback_type})"
                ai_status = f"Fallback ({fallback_type})"

            finding_num = len(findings) + 1
            print(f"[Situación H{finding_num}] Origen: {situation_source} | Longitud: {len(situation)}")

            if ai_result:
                proposals = [
                    clean_text(
                        proposal.proposal_text
                    )
                    for proposal
                    in ai_result.proposals
                    if clean_text(
                        proposal.proposal_text
                    )
                ]

                findings.append({
                    "title": (
                        clean_text(
                            ai_result.title
                        )
                        or hallazgo_text[:100]
                    ),
                    "source_number": finding_num,
                    "situation": situation,
                    "situation_source": situation_source,
                    "proposal": (
                        proposals[0]
                        if proposals
                        else proposal_text
                    ),
                    "proposals_ai": (
                        proposals
                    ),
                    "risk": (
                        clean_text(
                            ai_result.risk
                        )
                    ),
                    "severity": (
                        ai_result.severity
                        if ai_result.severity
                        in (
                            "Alto",
                            "Medio",
                            "Bajo"
                        )
                        else severity
                    ),
                    "responsible_area": (
                        clean_text(
                            ai_result
                            .responsible_area
                        )
                        or area
                    ),
                    "evidence": (
                        clean_text(
                            ai_result.evidence
                        )
                    ),
                    "cause": (
                        clean_text(
                            ai_result.cause
                        )
                    ),
                    "affected_process_or_control":
                        clean_text(
                            ai_result
                            .affected_process_or_control
                        ),
                    "impact": (
                        clean_text(
                            ai_result.impact
                        )
                    ),
                    "ai_confidence": (
                        ai_result.confidence
                    ),
                    "ai_validated": True,
                    "ai_status": ai_status,
                    "source_text": source_block,
                })

            else:
                findings.append({
                    "title": (
                        hallazgo_text[:100]
                    ),
                    "source_number": finding_num,
                    "situation": situation,
                    "situation_source": situation_source,
                    "proposal": proposal_text,
                    "proposals_ai": (
                        [proposal_text]
                        if proposal_text
                        else []
                    ),
                    "risk": "",
                    "severity": severity,
                    "responsible_area": area,
                    "evidence": "",
                    "cause": "",
                    "affected_process_or_control": "",
                    "impact": "",
                    "ai_confidence": 0,
                    "ai_validated": False,
                    "ai_status": ai_status,
                    "source_text": source_block,
                })

    return findings


def _extract_heuristic_paragraph_findings(raw_text, explicit_area="Operaciones"):
    lines = [clean_text(l) for l in raw_text.splitlines() if clean_text(l)]
    valid_lines = [l for l in lines if len(l) > 15 and not _is_non_finding_header(normalize_text(l))]

    if not valid_lines:
        return []

    findings = []
    chunk = []
    chunk_len = 0
    chunk_idx = 1

    for line in valid_lines:
        chunk.append(line)
        chunk_len += len(line)
        if chunk_len >= 300:
            full_block = " ".join(chunk)
            title = chunk[0][:90]
            findings.append({
                "title": title,
                "situation": full_block,
                "proposal": "",
                "proposals_ai": [],
                "risk": "",
                "severity": "Medio",
                "responsible_area": explicit_area,
                "evidence": "",
                "cause": "",
                "affected_process_or_control": "",
                "impact": "",
                "ai_confidence": 0,
                "ai_validated": False,
                "ai_status": "Fallback heurístico (Párrafos)",
                "source_text": full_block,
                "source_location": f"Párrafo / Bloque {chunk_idx}"
            })
            chunk_idx += 1
            chunk = []
            chunk_len = 0

    if chunk:
        full_block = " ".join(chunk)
        title = chunk[0][:90]
        findings.append({
            "title": title,
            "situation": full_block,
            "proposal": "",
            "proposals_ai": [],
            "risk": "",
            "severity": "Medio",
            "responsible_area": explicit_area,
            "evidence": "",
            "cause": "",
            "affected_process_or_control": "",
            "impact": "",
            "ai_confidence": 0,
            "ai_validated": False,
            "ai_status": "Fallback heurístico (Párrafos)",
            "source_text": full_block,
            "source_location": f"Párrafo / Bloque {chunk_idx}"
        })

    return findings[:10]


def _extract_findings_directly_via_ai(raw_text, filename):
    client = _get_openai_client()
    if not client:
        return []

    chunks = [raw_text[i:i+4000] for i in range(0, len(raw_text), 4000)]
    extracted = []

    for idx, chunk in enumerate(chunks[:5], start=1):
        ai_res = run_double_ai_review(
            source_text=chunk,
            fallback_title=f"Hallazgo Detectado por IA #{idx}",
            fallback_area="Operaciones",
            fallback_severity="Medio"
        )

        if ai_res and ai_res.situation:
            proposals = [clean_text(p.proposal_text) for p in ai_res.proposals if clean_text(p.proposal_text)]
            extracted.append({
                "title": clean_text(ai_res.title) or f"Hallazgo {idx}",
                "situation": clean_text(ai_res.situation),
                "risk": clean_text(ai_res.risk),
                "severity": ai_res.severity if ai_res.severity in ("Alto", "Medio", "Bajo") else "Medio",
                "responsible_area": clean_text(ai_res.responsible_area) or "Operaciones",
                "evidence": clean_text(ai_res.evidence),
                "cause": clean_text(ai_res.cause),
                "affected_process_or_control": clean_text(ai_res.affected_process_or_control),
                "impact": clean_text(ai_res.impact),
                "proposals_ai": proposals,
                "ai_confidence": ai_res.confidence,
                "ai_validated": True,
                "ai_status": "IA completa",
                "source_text": chunk,
                "source_location": f"Bloque de texto {idx}"
            })

    return extracted


def _has_explicit_finding_headers(raw_text):
    pattern = re.compile(
        r"(?:^|\n)\s*(?:\d+[\.\)]\s*)?(?:hallazgo|observaci[oó]n|desviaci[oó]n)\s*(?:n[°º]?\s*)?\d+",
        re.IGNORECASE
    )
    return bool(pattern.search(raw_text))


# ============================================================
# FUNCIÓN PRINCIPAL
# ============================================================

def parse_audit_report(
    file_path,
    filename
):
    raw_text = extract_raw_text_from_file(file_path)

    if not raw_text or not raw_text.strip():
        raise ValueError("No se pudo extraer texto del archivo o el contenido está vacío.")

    parsed = parse_document(
        raw_text,
        filename
    )

    extracted_findings = parsed.get("findings", [])

    if extracted_findings:
        print("[Parser] Modo hallazgos: explícito")
    elif _has_explicit_finding_headers(raw_text):
        print("[Parser] ERROR: se detectaron encabezados explícitos de hallazgos pero no pudieron extraerse.")
        extracted_findings = []
    else:
        if raw_text.strip():
            ai_direct_findings = _extract_findings_directly_via_ai(raw_text, filename)
            if ai_direct_findings:
                print("[Parser] Modo hallazgos: IA directa")
                extracted_findings = ai_direct_findings
            else:
                print("[Parser] Modo hallazgos: heurístico")
                extracted_findings = _extract_heuristic_paragraph_findings(raw_text, parsed.get("area", "Operaciones"))

    report_title = parsed.get(
        "report_title",
        f"Informe de Auditoría - {filename}"
    )

    explicit_area = parsed.get("area", "Operaciones")
    explicit_auditor = parsed.get("auditor", "Auditoría Interna")

    global_proposals_raw = parsed.get("global_proposals", [])
    parsed_proposals_list = []

    for gp in global_proposals_raw:
        p_num = gp["number"]
        p_text = clean_text(" ".join(gp.get("proposal_text_lines", [])))
        if not p_text:
            continue

        finding_numbers = _parse_finding_references(p_text)
        link_status = "linked" if finding_numbers else "unlinked"

        if finding_numbers:
            print(f"[Propuestas] Propuesta {p_num} detectada | refs: {finding_numbers}")
        else:
            print(f"[Propuestas] Propuesta {p_num} detectada | refs: []")

        parsed_proposals_list.append({
            "number": p_num,
            "proposal_text": p_text,
            "finding_numbers": finding_numbers,
            "link_status": link_status
        })

    if parsed_proposals_list:
        print(f"[Propuestas] Total detectadas: {len(parsed_proposals_list)}")

    relational_findings = []

    for index, finding in enumerate(extracted_findings, start=1):
        finding_id = str(uuid.uuid4())
        source_key = finding.get("source_key") or hashlib.md5(f"{filename}_{index}_{finding.get('title','')}".encode('utf-8')).hexdigest()[:12]
        source_item_id = finding.get("source_item_id") or f"src_{source_key}"
        source_item_ids = finding.get("source_item_ids") or [source_item_id]

        finding_number = finding.get("source_number", index)
        h_code = finding.get("code") or f"H-2026-{finding_number:03d}"

        title = clean_text(finding.get("title", f"Hallazgo {finding_number}"))
        situation = clean_text(finding.get("situation", title))
        if len(situation) > 500:
            situation = _summarize_text(situation, max_chars=500)
        risk = clean_text(finding.get("risk", ""))
        severity = finding.get("severity", "Medio")
        if severity not in ("Alto", "Medio", "Bajo"):
            severity = "Medio"

        area = clean_text(finding.get("responsible_area")) or explicit_area
        action_owner = clean_text(finding.get("action_owner")) or "Pendiente de definir"

        matching_global_proposals = [p for p in parsed_proposals_list if finding_number in p["finding_numbers"]]
        proposal_numbers = [p["number"] for p in matching_global_proposals]

        proposals = []
        if matching_global_proposals:
            for mp in matching_global_proposals:
                proposals.append({
                    "number": mp["number"],
                    "title": mp["proposal_text"],
                    "proposal_text": mp["proposal_text"],
                    "severity": severity,
                    "responsible_area": area,
                    "action_owner": action_owner,
                    "target_date": "",
                    "status": "Pendiente",
                    "action_plans": [],
                })
        else:
            proposal_texts = finding.get("proposals_ai") or []
            if not proposal_texts:
                legacy_proposal = clean_text(finding.get("proposal", ""))
                if legacy_proposal:
                    proposal_texts = [legacy_proposal]

            for proposal_text in proposal_texts:
                proposal_text = clean_text(proposal_text)
                if not proposal_text:
                    continue

                proposals.append({
                    "title": proposal_text,
                    "proposal_text": proposal_text,
                    "severity": severity,
                    "responsible_area": area,
                    "action_owner": action_owner,
                    "target_date": "",
                    "status": "Pendiente",
                    "action_plans": [],
                })

        relational_findings.append({
            "id": finding_id,
            "sourceItemId": source_item_id,
            "sourceItemIds": source_item_ids,
            "sourceKey": source_key,
            "sourceFile": filename,
            "sourceLocation": finding.get("source_location") or f"Sección / Fila {index}",
            "evidence": clean_text(finding.get("evidence", "")),
            "included": True,
            "selectedAsFinding": True,
            "converted": False,
            "code": h_code,
            "number": index,
            "source_number": finding_number,
            "title": title,
            "situation": situation,
            "risk": risk,
            "severity": severity,
            "responsible_area": area,
            "action_owner": action_owner,
            "status": "Pendiente",
            "observations": finding.get("observations", ""),
            "proposal_numbers": proposal_numbers,
            "proposals": proposals,
            "cause": clean_text(finding.get("cause", "")),
            "affected_process_or_control": clean_text(finding.get("affected_process_or_control", "")),
            "impact": clean_text(finding.get("impact", "")),
            "ai_confidence": finding.get("ai_confidence", 0),
            "ai_validated": finding.get("ai_validated", False),
            "ai_status": finding.get("ai_status") or ("IA completa" if finding.get("ai_validated") else "Fallback heurístico")
        })

    unlinked_count = sum(1 for p in parsed_proposals_list if p["link_status"] == "unlinked")
    print(f"[Parser] Hallazgos: {len(relational_findings)} | Propuestas: {len(parsed_proposals_list)} | Propuestas sin vincular: {unlinked_count}")

    return {
        "report": {
            "title": report_title,
            "process": explicit_area or "Control Interno",
            "area": explicit_area,
            "period": "2026",
            "auditor": explicit_auditor,
            "summary": f"Informe {filename} procesado con {len(relational_findings)} hallazgos y {len(parsed_proposals_list)} propuestas.",
            "source_filename": filename
        },
        "findings": relational_findings,
        "proposals": parsed_proposals_list
    }
