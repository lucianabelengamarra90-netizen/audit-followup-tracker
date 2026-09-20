import unittest
import json
import os
import sys
import tempfile
import docx
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

_TMP_PARSER_PATH = tempfile.mktemp(suffix=".db")
os.environ["DATABASE_PATH"] = _TMP_PARSER_PATH

import database
from app import app
from report_parser import parse_audit_report, extract_raw_text_from_file


class ParserRegressionTests(unittest.TestCase):
    """
    Suite congelada de pruebas de regresión para garantizar la inmutabilidad
    del contrato y los resultados de report_parser.py y la ruta /parse-preview.
    
    Nota sobre IA: La clave OPENAI_API_KEY se fuerza a vacía ("") en setUp()
    para garantizar la ejecución determinística del flujo heurístico fallback
    sin depender ni realizar llamadas a APIs externas.
    """

    def setUp(self):
        self.tmp_fd, self.tmp_path = tempfile.mkstemp(suffix=".db")
        os.close(self.tmp_fd)
        os.environ["DATABASE_PATH"] = self.tmp_path
        database.DB_PATH = self.tmp_path
        database.init_db()
        self.app = app.test_client()
        self.app.testing = True
        os.environ["OPENAI_API_KEY"] = ""

    def tearDown(self):
        if hasattr(self, "tmp_path") and os.path.exists(self.tmp_path):
            try:
                os.remove(self.tmp_path)
            except Exception:
                pass

    def test_txt_heuristic_report_parsing_structure(self):
        txt_content = """INFORME DE AUDITORÍA INTERNA 2026

1. HALLAZGOS DETECTADOS

Hallazgo 1: Inexistencia de doble firma en transferencias bancarias
Situación Observada: Se detectaron pagos realizados sin la autorización conjunta.
Riesgo: Alto
Área: Tesorería

7. PROPUESTAS DE MEJORA
Propuesta 1 (relacionada con Hallazgo 1): Configurar firma conjunta obligatoria en el portal bancario.
"""
        with tempfile.NamedTemporaryFile(suffix=".txt", delete=False, mode="w", encoding="utf-8") as f:
            f.write(txt_content)
            tmp_path = f.name

        try:
            parsed = parse_audit_report(tmp_path, "informe_tesoreria.txt")
            
            # Verificación del contrato estricto devuelto por parse_audit_report
            self.assertIn("report", parsed)
            self.assertIn("findings", parsed)
            self.assertIn("proposals", parsed)

            report = parsed["report"]
            self.assertEqual(report["title"], "INFORME DE AUDITORÍA INTERNA 2026")
            
            findings = parsed["findings"]
            self.assertEqual(len(findings), 1)
            f1 = findings[0]
            self.assertEqual(f1["source_number"], 1)
            self.assertEqual(f1["severity"], "Alto")
            self.assertIn("Inexistencia de doble firma", f1["title"])
            
            # Verificación de propuestas vinculadas
            proposals = f1["proposals"]
            self.assertEqual(len(proposals), 1)
            self.assertIn("Configurar firma conjunta", proposals[0]["proposal_text"])
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)

    def test_csv_deterministic_report_parsing(self):
        csv_content = "Hallazgo;Propuesta;Area;Riesgo\nFalta conciliación de inventario;Realizar conteos físicos trimestrales;Logística;Medio\n"
        with tempfile.NamedTemporaryFile(suffix=".csv", delete=False, mode="w", encoding="utf-8") as f:
            f.write(csv_content)
            tmp_path = f.name

        try:
            parsed = parse_audit_report(tmp_path, "inventario.csv")
            findings = parsed.get("findings", [])
            self.assertEqual(len(findings), 1)
            f1 = findings[0]
            self.assertIn("Falta conciliación", f1["title"])
            self.assertEqual(f1["severity"], "Medio")
            self.assertEqual(len(f1["proposals"]), 1)
            self.assertIn("Realizar conteos físicos", f1["proposals"][0]["proposal_text"])
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)

    def test_pdf_deterministic_report_parsing(self):
        pdf_bytes = (
            b'%PDF-1.4\n'
            b'1 0 obj <</Type /Catalog /Pages 2 0 R>> endobj\n'
            b'2 0 obj <</Type /Pages /Kinds [] /Count 1 /Kids [3 0 R]>> endobj\n'
            b'3 0 obj <</Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources <</Font <</F1 5 0 R>>>> >> endobj\n'
            b'4 0 obj <</Length 120>> stream\n'
            b'BT /F1 12 Tf 50 700 Td (INFORME AUDITORIA PDF 2026) Tj ET\n'
            b'BT /F1 12 Tf 50 670 Td (Hallazgo 1: Control de acceso ineficiente en servidores) Tj ET\n'
            b'BT /F1 12 Tf 50 640 Td (Propuesta 1: Restringir accesos a usuarios administradores) Tj ET\n'
            b'endstream\n'
            b'endobj\n'
            b'5 0 obj <</Type /Font /Subtype /Type1 /BaseFont /Helvetica>> endobj\n'
            b'xr\n'
            b'0 6\n'
            b'0000000000 65535 f \n'
            b'0000000009 00000 n \n'
            b'0000000056 00000 n \n'
            b'0000000125 00000 n \n'
            b'0000000245 00000 n \n'
            b'0000000414 00000 n \n'
            b'trailer <</Size 6 /Root 1 0 R>>\n'
            b'startxref\n'
            b'489\n'
            b'%%EOF\n'
        )
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
            f.write(pdf_bytes)
            tmp_path = f.name

        try:
            parsed = parse_audit_report(tmp_path, "reporte_seguridad.pdf")
            findings = parsed.get("findings", [])
            self.assertEqual(len(findings), 1)
            self.assertIn("Control de acceso", findings[0]["title"])
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)

    def test_docx_deterministic_report_parsing(self):
        doc = docx.Document()
        doc.add_heading('Informe de Auditoría DOCX 2026', level=1)
        doc.add_paragraph('Hallazgo 1: Falta de backups automatizados en servidores')
        doc.add_paragraph('Propuesta 1: Implementar sistema de respaldos diarios')
        
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            doc.save(f.name)
            tmp_path = f.name

        try:
            parsed = parse_audit_report(tmp_path, "reporte_backups.docx")
            findings = parsed.get("findings", [])
            self.assertEqual(len(findings), 1)
            self.assertIn("Falta de backups", findings[0]["title"])
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)

    def test_xlsx_deterministic_report_parsing(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Hallazgos y Propuestas'
        headers = [
            'Área / Proceso', 'ID Hallazgo', 'Hallazgo (Situación Observada)',
            'ID Propuesta', 'Propuesta de Mejora (Recomendación)', 'Riesgo',
            'Responsable', 'Fecha compromiso', 'Estado', 'Avance', 'Acciones', 'Observaciones'
        ]
        ws.append(headers)
        ws.append([
            'Sistemas', 'H-001', 'Sin redundancia de energía en datacenter',
            'PM-001', 'Instalar equipo UPS industrial', 'Alto', 'Operaciones',
            '2026-10-31', 'En proceso', '0%', '', ''
        ])

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            wb.save(f.name)
            tmp_path = f.name

        try:
            parsed = parse_audit_report(tmp_path, "planilla_auditoria.xlsx")
            findings = parsed.get("findings", [])
            self.assertGreater(len(findings), 0)
            has_matching_situation = any("Sin redundancia" in (f.get("situation") or "") or "Sin redundancia" in (f.get("title") or "") for f in findings)
            self.assertTrue(has_matching_situation)
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)

    def test_parse_preview_endpoint_contract(self):
        txt_content = """Informe de Prueba
Hallazgo 1: Falta de respaldo de servidores
Propuesta 1: Implementar respaldos diarios automatizados en la nube.
"""
        with tempfile.NamedTemporaryFile(suffix=".txt", delete=False, mode="w", encoding="utf-8") as f:
            f.write(txt_content)
            tmp_path = f.name

        try:
            with open(tmp_path, "rb") as f_bytes:
                data = {'file': (f_bytes, "test_preview.txt")}
                res = self.app.post("/parse-preview", data=data, content_type='multipart/form-data')
                
            self.assertEqual(res.status_code, 200)
            json_data = json.loads(res.data)
            
            # Verificación estricta del contrato devuelto por el endpoint HTTP /parse-preview
            self.assertIn("report", json_data)
            self.assertIn("findings", json_data)
            self.assertIn("report_file", json_data)
            self.assertEqual(json_data["report_file"], "test_preview.txt")

            findings = json_data["findings"]
            self.assertGreater(len(findings), 0)
            self.assertIn("code", findings[0])
            self.assertIn("situation", findings[0])
            self.assertIn("proposals", findings[0])
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)


if __name__ == "__main__":
    unittest.main()
