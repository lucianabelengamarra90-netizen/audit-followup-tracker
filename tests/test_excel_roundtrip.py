import unittest
import os
import tempfile
import json
from io import BytesIO
from openpyxl import load_workbook

_TMP_EXCEL_PATH = tempfile.mktemp(suffix=".db")
os.environ["DATABASE_PATH"] = _TMP_EXCEL_PATH

import database
from app import app


class ExcelRoundTripTestCase(unittest.TestCase):

    def setUp(self):
        self.db_fd, self.db_path = tempfile.mkstemp(suffix=".db")
        os.close(self.db_fd)
        database.DB_PATH = self.db_path
        app.config["TESTING"] = True
        self.client = app.test_client()
        database.init_db()

    def tearDown(self):
        if os.path.exists(self.db_path):
            os.remove(self.db_path)

    def test_excel_export_import_roundtrip_integrity(self):
        # 1. Ingest sample relational report
        report_data = {
            "title": "Informe de Prueba Roundtrip Excel",
            "process": "Proceso Tesorería",
            "area": "Finanzas",
            "period": "2026",
            "auditor": "Auditoría Interna",
            "summary": "Verificación de exportación e importación sin pérdida de datos."
        }
        findings_hierarchy = [
            {
                "code": "H-2026-901",
                "title": "Diferencia de arqueo",
                "situation": "Arqueo de caja chica con faltante no justificado",
                "risk": "Faltante de valores",
                "severity": "Alto",
                "responsible_area": "Tesorería",
                "action_owner": "Carlos Pérez",
                "status": "En proceso",
                "observations": "Requiere regularización",
                "proposals": [
                    {
                        "code": "PM-2026-901",
                        "title": "Arqueos sorpresivos semanal",
                        "proposal_text": "Implementar protocolo de arqueos sorpresivos semanalmente",
                        "target_date": "2026-11-30",
                        "status": "En proceso",
                        "action_plans": [
                            {
                                "code": "PA-2026-901",
                                "title": "Redactar instructivo de arqueos",
                                "action_text": "Redactar instructivo de arqueos y notificar al personal",
                                "action_owner": "Carlos Pérez",
                                "target_date": "2026-11-15",
                                "progress_pct": 50,
                                "status": "En proceso",
                                "notes": "Avance en borrador"
                            }
                        ]
                    }
                ]
            }
        ]

        rep_id = database.save_relational_report_structure(report_data, findings_hierarchy, "informe_test.xlsx")
        self.assertTrue(rep_id)

        # 2. Export to Excel
        response_export = self.client.get("/export-excel")
        self.assertEqual(response_export.status_code, 200)
        self.assertEqual(response_export.mimetype, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        excel_bytes = response_export.data
        self.assertGreater(len(excel_bytes), 0)

        # Validate openpyxl read of exported Excel
        wb = load_workbook(filename=BytesIO(excel_bytes))
        self.assertIn("Hallazgos y Propuestas", wb.sheetnames)
        self.assertIn("Propuestas de Mejora", wb.sheetnames)
        self.assertIn("Planes de Acción", wb.sheetnames)

        # 3. Import back into clean DB
        database.DB_PATH = tempfile.mktemp(suffix=".db")
        database.init_db()

        data_stream = BytesIO(excel_bytes)
        response_import = self.client.post(
            "/import-excel",
            content_type="multipart/form-data",
            data={"file": (data_stream, "Reporte_AuditTrack_Roundtrip.xlsx")}
        )

        self.assertEqual(response_import.status_code, 200)
        json_resp = json.loads(response_import.data)
        self.assertTrue(json_resp.get("success"))
        self.assertEqual(json_resp.get("imported_count"), 1)

        # 4. Assert Imported Integrity
        findings = database.get_all_findings()
        self.assertEqual(len(findings), 1)
        imported_f = findings[0]
        self.assertIn("Arqueo de caja chica", imported_f["situation"])
        self.assertEqual(imported_f["severity"], "Alto")

        proposals = database.get_all_proposals()
        self.assertEqual(len(proposals), 1)
        self.assertIn("arqueos sorpresivos", proposals[0]["proposal_text"].lower())

        action_plans = database.get_all_action_plans()
        self.assertEqual(len(action_plans), 1)
        self.assertEqual(action_plans[0]["progress_pct"], 50)


if __name__ == "__main__":
    unittest.main()
