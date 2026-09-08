from flask import Flask, render_template, request, jsonify, send_file
import os
import re
from datetime import datetime
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from database import (
    init_db,
    save_report_and_findings,
    get_all_reports,
    get_report,
    get_findings,
    get_finding,
    update_finding_status,
    delete_finding,
    delete_report,
    get_dashboard_kpis
)
from report_parser import parse_audit_report, clean_text

init_db()

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 250 * 1024 * 1024
UPLOAD_FOLDER = os.path.join(os.path.dirname(os.path.abspath(__file__)), "uploads")
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER

ALLOWED_EXTENSIONS = {"xlsx", "csv", "docx", "pdf", "txt"}


def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/health")
def health():
    api_key = os.getenv("OPENAI_API_KEY", "").strip()
    return jsonify({
        "status": "ok",
        "app": "Audit Follow-up Tracker",
        "openaiConfigured": bool(api_key)
    })


@app.route("/upload-report", methods=["POST"])
def upload_report():
    if "file" not in request.files:
        return jsonify({"error": "No se seleccionó ningún archivo de informe."}), 400

    file = request.files["file"]
    if not file or not file.filename:
        return jsonify({"error": "El archivo enviado no es válido."}), 400

    if not allowed_file(file.filename):
        return jsonify({"error": "Formato de archivo no soportado. Usá PDF, Word (.docx) o Excel (.xlsx)."}), 400

    safe_filename = clean_text(file.filename).replace(" ", "_")
    saved_path = os.path.join(app.config["UPLOAD_FOLDER"], safe_filename)
    file.save(saved_path)

    try:
        parsed_data = parse_audit_report(saved_path, safe_filename)
        report_info = parsed_data.get("report", {})
        findings_info = parsed_data.get("findings", [])

        report_id, count = save_report_and_findings(report_info, findings_info, safe_filename)

        return jsonify({
            "message": f"Informe '{report_info.get('title')}' ingresado correctamente con {count} hallazgos/mejoras.",
            "report_id": report_id,
            "findings_count": count,
            "report": report_info,
            "findings": findings_info
        })
    except Exception as exc:
        print(f"Error procesando informe: {exc}")
        return jsonify({"error": f"No se pudo procesar el informe: {str(exc)}"}), 500


@app.route("/reports", methods=["GET"])
def list_reports():
    reports = get_all_reports()
    return jsonify({"reports": reports, "count": len(reports)})


@app.route("/reports/<report_id>", methods=["GET", "DELETE"])
def report_detail(report_id):
    if request.method == "DELETE":
        deleted = delete_report(report_id)
        if deleted:
            return jsonify({"message": "Informe y sus hallazgos eliminados correctamente."})
        return jsonify({"error": "Informe no encontrado."}), 404

    report = get_report(report_id)
    if not report:
        return jsonify({"error": "Informe no encontrado."}), 404
    return jsonify(report)


@app.route("/findings", methods=["GET"])
def list_findings():
    filters = {
        "status": request.args.get("status"),
        "severity": request.args.get("severity"),
        "process": request.args.get("process"),
        "area": request.args.get("area"),
        "search": request.args.get("search")
    }
    findings = get_findings(filters)
    return jsonify({"findings": findings, "count": len(findings)})


@app.route("/findings/<finding_id>", methods=["GET", "DELETE"])
def finding_detail(finding_id):
    if request.method == "DELETE":
        deleted = delete_finding(finding_id)
        if deleted:
            return jsonify({"message": "Hallazgo eliminado."})
        return jsonify({"error": "Hallazgo no encontrado."}), 404

    finding = get_finding(finding_id)
    if not finding:
        return jsonify({"error": "Hallazgo no encontrado."}), 404
    return jsonify(finding)


@app.route("/findings/<finding_id>/status", methods=["POST"])
def update_status(finding_id):
    data = request.get_json(silent=True) or {}
    new_status = clean_text(data.get("status", ""))
    notes = data.get("notes")
    target_date = data.get("target_date")
    evidence_file = data.get("evidence_file")

    if not new_status:
        return jsonify({"error": "El nuevo estado es obligatorio."}), 400

    updated = update_finding_status(finding_id, new_status, notes, target_date, evidence_file)
    if updated:
        return jsonify({"message": f"Estado actualizado a '{new_status}'."})
    return jsonify({"error": "No se encontró el hallazgo."}), 404


@app.route("/dashboard-stats")
def dashboard_stats():
    stats = get_dashboard_kpis()
    return jsonify(stats)


@app.route("/export-excel", methods=["POST"])
def export_excel():
    findings = get_findings()
    stats = get_dashboard_kpis()
    reports = get_all_reports()

    wb = Workbook()
    ws_kpi = wb.active
    ws_kpi.title = "Tablero de Control"
    ws_kpi.sheet_properties.tabColor = "17365D"

    # Estilos
    NAVY = "17365D"
    BLUE = "1F4E78"
    LIGHT_BLUE = "D9EAF7"
    TEXT_COLOR = "1F2937"
    BORDER_COLOR = "D0D7DE"
    THIN_BORDER = Border(
        left=Side(style="thin", color=BORDER_COLOR),
        right=Side(style="thin", color=BORDER_COLOR),
        top=Side(style="thin", color=BORDER_COLOR),
        bottom=Side(style="thin", color=BORDER_COLOR)
    )

    # Título
    ws_kpi.merge_cells("A1:H1")
    title_cell = ws_kpi["A1"]
    title_cell.value = "REPORTE CONSOLIDADO DE SEGUIMIENTO DE AUDITORÍA INTERNA"
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill("solid", fgColor=NAVY)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_kpi.row_dimensions[1].height = 40

    # KPIs resumen
    kpi_items = [
        ("Informes Evaluados", stats.get("total_reports", 0)),
        ("Total Recomendaciones", stats.get("total_findings", 0)),
        ("Implementados / Cerrados", stats.get("closed_findings", 0)),
        ("En Proceso", stats.get("in_progress_findings", 0)),
        ("En Revisión", stats.get("in_review_findings", 0)),
        ("Vencidos", stats.get("overdue_findings", 0)),
        ("% Cumplimiento Global", f"{stats.get('resolution_rate', 0)}%"),
    ]

    for col_idx, (label, val) in enumerate(kpi_items, start=1):
        cell_lbl = ws_kpi.cell(row=3, column=col_idx, value=label)
        cell_lbl.font = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
        cell_lbl.fill = PatternFill("solid", fgColor=BLUE)
        cell_lbl.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        cell_val = ws_kpi.cell(row=4, column=col_idx, value=val)
        cell_val.font = Font(name="Calibri", size=14, bold=True, color=TEXT_COLOR)
        cell_val.alignment = Alignment(horizontal="center", vertical="center")
        cell_val.border = THIN_BORDER

    ws_kpi.row_dimensions[3].height = 25
    ws_kpi.row_dimensions[4].height = 30

    # Solapa Detalle de Recomendaciones
    ws_detail = wb.create_sheet("Seguimiento Detallado")
    ws_detail.sheet_properties.tabColor = "5B9BD5"

    headers = [
        "N°", "Código", "Informe", "Proceso / Subproceso", "Oportunidad de Mejora / Hallazgo",
        "Situación Observada", "Riesgo / Impacto", "Propuesta de Acción",
        "Criticidad", "Área Responsable", "Responsable Plan", "Fecha Compromiso",
        "Estado", "Notas de Seguimiento", "Evidencia"
    ]

    for col_idx, h in enumerate(headers, start=1):
        c = ws_detail.cell(row=1, column=col_idx, value=h)
        c.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = THIN_BORDER
    ws_detail.row_dimensions[1].height = 30

    for idx, f in enumerate(findings, start=1):
        row_idx = idx + 1
        vals = [
            idx, f.get("code", ""), f.get("report_title", ""), f.get("process_step", ""),
            f.get("title", ""), f.get("situation", ""), f.get("risk", ""), f.get("proposal", ""),
            f.get("severity", ""), f.get("responsible_area", ""), f.get("action_owner", ""),
            f.get("target_date", ""), f.get("status", ""), f.get("follow_up_notes", ""),
            f.get("evidence_file", "")
        ]
        for col_idx, v in enumerate(vals, start=1):
            c = ws_detail.cell(row=row_idx, column=col_idx, value=v)
            c.font = Font(name="Calibri", size=9, color=TEXT_COLOR)
            c.border = THIN_BORDER
            c.alignment = Alignment(vertical="center", wrap_text=True)
            if col_idx in (1, 2, 9, 12, 13):
                c.alignment = Alignment(horizontal="center", vertical="center")

        ws_detail.row_dimensions[row_idx].height = 35

    # Auto ajustar anchos
    widths = [6, 12, 28, 25, 30, 45, 35, 38, 12, 20, 20, 16, 18, 30, 20]
    for idx, w in enumerate(widths, start=1):
        col_letter = get_column_letter(idx)
        ws_detail.column_dimensions[col_letter].width = w
        ws_kpi.column_dimensions[col_letter].width = 18

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Seguimiento_Auditoria_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5002))
    host = os.environ.get("HOST", "127.0.0.1")
    app.run(host=host, port=port, debug=False)
